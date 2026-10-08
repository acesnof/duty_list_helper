from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
OUT=Path(__file__).parent
D=Document(); s=D.sections[0]
s.page_width=Pt(595.28); s.page_height=Pt(841.89)
s.top_margin=Pt(29); s.bottom_margin=Pt(25); s.left_margin=Pt(59); s.right_margin=Pt(49)
s.header_distance=Pt(0); s.footer_distance=Pt(12)
style=D.styles['Normal']; style.font.name='Calibri'; style.font.size=Pt(12)
style.paragraph_format.space_after=Pt(0); style.paragraph_format.line_spacing=Pt(15)
style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)

def pf(p,size=12,bold=False,align=None,line=15):
 p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=Pt(line)
 if align is not None:p.alignment=align
 for r in p.runs:r.font.name='Calibri';r.font.size=Pt(size);r.bold=bold
 return p

def para(txt='',size=12,bold=False,height=15,align=None,before=0,left=0):
 p=D.add_paragraph(txt);pf(p,size,bold,align,height);p.paragraph_format.space_before=Pt(before);p.paragraph_format.left_indent=Pt(left);return p

def gap(h):return para('',height=h,size=1)

def runs(p,parts):
 for text,b in parts:
  r=p.add_run(text);r.bold=b
 return p

def borders(cell,values):
 pr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:tcBorders')
 for edge,val in values.items():
  a=OxmlElement('w:'+edge)
  for k,v in val.items(): a.set(qn('w:'+k),str(v))
  e.append(a)
 pr.append(e)

def table(widths,heights,border=False):
 t=D.add_table(rows=len(heights),cols=len(widths)); t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.LEFT
 if border:
  ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'80');ind.set(qn('w:type'),'dxa');t._tbl.tblPr.append(ind)
 for c,w in zip(t.columns,widths):c.width=Pt(w)
 for row,h in zip(t.rows,heights):
  row.height=Pt(h);row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
  for c,w in zip(row.cells,widths):
   c.width=Pt(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
   for k,n in [('top',0),('bottom',0),('left',4 if border else 0),('right',2 if border else 0)]:
    a=OxmlElement('w:'+k);a.set(qn('w:w'),str(n*20));a.set(qn('w:type'),'dxa');mar.append(a)
   pr.append(mar);pf(c.paragraphs[0])
   borders(c,{e:{'val':'single' if border else 'nil','sz':6,'color':'000000'} for e in ['top','bottom','left','right']})
 return t

def cell(c,txt,bold=False,size=12,align=None,before=0):
 p=c.paragraphs[0];p.text=txt;pf(p,size,bold,align);p.paragraph_format.space_before=Pt(before);return p

def add(c,txt,bold=False,size=12,align=None,before=0):
 p=c.add_paragraph(txt);pf(p,size,bold,align);p.paragraph_format.space_before=Pt(before);return p
C=WD_ALIGN_PARAGRAPH.CENTER
para('APPLICATION FORM FOR LEAVE',16,True,21,C)
para('(SUBMIT IN ONE COPY)',12,True,17,C)
gap(6)
t=table([110,132,106,64,74],[51,18,21,26],True)
a=t.cell(0,0).merge(t.cell(1,0));cell(a,'BRANCH/PILLAR',True);add(a,'LDG',align=C,before=14)
for j,label,val in [(1,'1. Name and surname','Madeira Roberto'),(2,'2. Grade/Rank','OR8')]:
 cell(t.cell(0,j),label,True);add(t.cell(0,j),val,align=C)
a=t.cell(0,3).merge(t.cell(0,4));cell(a,'3. Nationality and Service',True,size=11.5);add(a,'PRT - ARMY',align=C)
cell(t.cell(2,0),'POSITION NUMBER',True,size=11.5)
cell(t.cell(3,0),'100/20',align=C)
for j,label,val in [(1,'4. Leave starts','04/11/2026'),(2,'5. Leave ends','24/11/2026')]:
 a=t.cell(1,j).merge(t.cell(3,j));cell(a,label,True);add(a,val,align=C,before=10)
a=t.cell(1,3).merge(t.cell(1,4));cell(a,'6. No. of days requested',True,size=11.5)
cell(t.cell(2,3),'OUT OF THE\nMISSION',True,8,C); t.cell(2,3).paragraphs[0].paragraph_format.line_spacing=Pt(9)
cell(t.cell(2,4),'DAYS OF LEAVE',True,8,C,before=5)
cell(t.cell(3,3),'21',align=C,before=3);cell(t.cell(3,4),'11',align=C,before=3)
gap(15)
t=table([109,111,28,110,28,100],[16,16])
cell(t.cell(0,0),'7. TYPE OF LEAVE:',True)
for row,left,right,mark in [(0,'Ordinary Leave','Short time Leave','☐'),(1,'National Holidays','Others','☒')]:
 cell(t.cell(row,1),left); p=cell(t.cell(row,2),mark,size=17);p.runs[0].font.name='Segoe UI Symbol'
 cell(t.cell(row,3),right);p=cell(t.cell(row,4),'☐',size=17);p.runs[0].font.name='Segoe UI Symbol'
gap(24)
para('8. LEAVE ADDRESS: Rua da Bela Vista nº31 1ºandar',bold=True)
gap(15)
p=para('');runs(p,[('9. Telephone number of contact: ',True),('+351 967592231',False)])
gap(15)
para('I have the sufficient leave to cover the above period of ordinary leave requested and I have not been\nappointed to any special duties or assignments requiring my presence at any time during the period\nrequested.',size=11.5,height=15)
gap(29)
t=table([111,187,156,32],[1,16])
for col in [0,2]:
 borders(t.cell(0,col),{'bottom':{'val':'single','sz':6,'color':'000000'}})
cell(t.cell(1,0),'(Date)',align=C)
cell(t.cell(1,2),'(Signature of the applicant)',align=C)
gap(8)
p=para('',height=4,size=1)
ppr=p._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom');b.set(qn('w:val'),'triple');b.set(qn('w:sz'),'9');b.set(qn('w:space'),'0');pb.append(b);ppr.append(pb)
gap(8)
para('APPROVAL:',bold=True)
gap(15)
p=para('');runs(p,[('10. Verification by J1: ',True),('The request is according to paragraph (2) (b).  ',False),('________________________',False)])
gap(15)
p=para('');runs(p,[('11. COORDINATION 1:  ',True),('Branch/Pillar Chief: _______________________',False)])
gap(15)
p=para('',left=36);runs(p,[('POC NOMINATED BY CHIEF DURING ABSENCE: ',True),('OF 1 Ruivaco Liane',True)]);p.runs[-1].underline=True
gap(15)
para('SIGNATURE OF POC TAKING KNOWLODGE:  _______________________________________',left=36)
gap(15)
p=para('',left=36);runs(p,[('REASON FOR DEVIATION FROM LEAVE PLAN (If any): ',True),('None_____________________',False)])
gap(15)
p=para('');runs(p,[('12. COORDINATION 2: ',True),('Senior National Representative: _________________________________',False)])
para('Senior national representative will verify that the leave is taken in accordance with National Regulations.',size=9.5,bold=True,height=13)
gap(28)
para('13.  COS APPROVAL:',bold=True)
gap(15)
p=para('___________________________________',align=C,left=17)
f=s.footer.paragraphs[0];f.text='UNCLASSIFIED';pf(f,10,True,C,17)
f=s.footer.add_paragraph('RELEASABLE TO EUTM RCA TCN');pf(f,9,False,C,16)
f=s.footer.add_paragraph('Annex A  1 of  1');pf(f,10,False,WD_ALIGN_PARAGRAPH.RIGHT,12)
D.core_properties.title='Application Form for Leave';D.core_properties.author=''
D.save(OUT/'Application Form For Leave.docx')
print(OUT/'Application Form For Leave.docx')
